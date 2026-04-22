import os
import pandas as pd
from io import BytesIO

from openpyxl.styles import Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

REQUEST_FILE = "data/shift_requests.csv"


def load_requests():
    """
    スタッフの希望シフトCSVを読み込む
    なければ空のDataFrameを返す
    """
    if os.path.exists(REQUEST_FILE):
        return pd.read_csv(REQUEST_FILE, encoding="utf-8-sig")
    else:
        return pd.DataFrame(columns=["名前", "昼", "夜"])


def save_requests(df):
    """
    スタッフの希望シフトCSVを保存する
    """
    df.to_csv(REQUEST_FILE, index=False, encoding="utf-8-sig")


def upsert_request(name, lunch_days, night_days):
    """
    スタッフの希望シフトを追加または更新する
    """
    df = load_requests()

    lunch_text = ",".join(map(str, sorted(lunch_days)))
    night_text = ",".join(map(str, sorted(night_days)))

    if name in df["名前"].values:
        df.loc[df["名前"] == name, "昼"] = lunch_text
        df.loc[df["名前"] == name, "夜"] = night_text
    else:
        new_row = pd.DataFrame({
            "名前": [name],
            "昼": [lunch_text],
            "夜": [night_text]
        })
        df = pd.concat([df, new_row], ignore_index=True)

    df = df.sort_values("名前").reset_index(drop=True)
    save_requests(df)
    return df


def text_to_days(text):
    """
    '1,3,5' のような文字列を [1,3,5] に変換する
    """
    if pd.isna(text) or str(text).strip() == "":
        return []

    return list(map(int, filter(None, str(text).replace(".0", "").split(","))))


def create_monthly_summary(df):
    """
    日ごとの昼・夜メンバー一覧を作る
    """
    days = list(range(1, 32))
    shift_table = []

    for day in days:
        lunch_members = []
        night_members = []

        for _, row in df.iterrows():
            if pd.notna(row["昼"]):
                lunch_days = text_to_days(row["昼"])
                if day in lunch_days:
                    lunch_members.append(row["名前"])

            if pd.notna(row["夜"]):
                night_days = text_to_days(row["夜"])
                if day in night_days:
                    night_members.append(row["名前"])

        shift_table.append({
            "日": day,
            "昼メンバー": ",".join(lunch_members),
            "昼人数": len(lunch_members),
            "夜メンバー": ",".join(night_members),
            "夜人数": len(night_members)
        })

    return pd.DataFrame(shift_table)


def build_output_table(long_df):
    """
    最終出力用の表を作る
    行：昼・夜
    列：1日〜31日
    """
    days = list(range(1, 32))
    rows = []

    for kubun in ["昼", "夜"]:
        row_data = {"区分": kubun}

        for day in days:
            target = long_df[
                (long_df["日"] == day) &
                (long_df["区分"] == kubun)
            ]

            if target.empty:
                row_data[f"{day}日"] = ""
            else:
                target = target.sort_values(["開始時間", "名前"])

                names = []
                for _, r in target.iterrows():
                    name = r["名前"]
                    time = r["開始時間"]

                    if bool(r.get("厨房担当", False)):
                        names.append(f"{name}(厨) {time}")
                    else:
                        names.append(f"{name} {time}")

                row_data[f"{day}日"] = " / ".join(names)

        rows.append(row_data)

    return pd.DataFrame(rows)


def to_excel_bytes(df, long_df):
    """
    Excel出力
    - 1週間ごとに縦表示
    - 厨房担当の人だけ赤文字
    - セル幅、行の高さを調整
    """
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book
        ws = workbook.create_sheet("シフト")

        # 最初に自動作成される空シートを削除
        if "Sheet" in workbook.sheetnames:
            std = workbook["Sheet"]
            workbook.remove(std)

        thin = Side(style="thin")

        # 1週間ごとに縦並び
        week_ranges = [
            range(1, 8),
            range(8, 15),
            range(15, 22),
            range(22, 29),
            range(29, 32)
        ]

        start_row = 1

        for week_index, week_days in enumerate(week_ranges, start=1):
            # 週タイトル
            ws.cell(row=start_row, column=1, value=f"{week_index}週目")

            # ヘッダー
            ws.cell(row=start_row + 1, column=1, value="区分")
            for col_idx, day in enumerate(week_days, start=2):
                ws.cell(row=start_row + 1, column=col_idx, value=f"{day}日")

            # 昼・夜
            for row_offset, kubun in enumerate(["昼", "夜"], start=2):
                current_row = start_row + row_offset
                ws.cell(row=current_row, column=1, value=kubun)

                for col_idx, day in enumerate(week_days, start=2):
                    target = long_df[
                        (long_df["日"] == day) &
                        (long_df["区分"] == kubun)
                    ]

                    cell = ws.cell(row=current_row, column=col_idx)

                    if target.empty:
                        cell.value = ""
                    else:
                        target = target.sort_values(["開始時間", "名前"])

                        rich_parts = []
                        red_font = InlineFont(color="00FF0000", b=True)

                        for idx, (_, r) in enumerate(target.iterrows()):
                            name = str(r["名前"])
                            time = str(r["開始時間"])
                            is_kitchen = bool(r.get("厨房担当", False))

                            text = f"{name}(厨) {time}" if is_kitchen else f"{name} {time}"

                            if is_kitchen:
                                rich_parts.append(TextBlock(red_font, text))
                            else:
                                rich_parts.append(text)

                            if idx < len(target) - 1:
                                rich_parts.append("\r\n")

                        cell.value = CellRichText(*rich_parts)

            # 罫線・中央揃え
            max_col = len(week_days) + 1
            for row in ws.iter_rows(
                min_row=start_row,
                max_row=start_row + 3,
                min_col=1,
                max_col=max_col
            ):
                for cell in row:
                    cell.alignment = Alignment(
                        wrap_text=True,
                        vertical="center",
                        horizontal="center"
                    )
                    cell.border = Border(
                        left=thin,
                        right=thin,
                        top=thin,
                        bottom=thin
                    )

            # 列幅
            ws.column_dimensions["A"].width = 10
            for col_idx in range(2, max_col + 1):
                col_letter = ws.cell(row=start_row + 1, column=col_idx).column_letter
                ws.column_dimensions[col_letter].width = 20

            # 行の高さ
            ws.row_dimensions[start_row].height = 22
            ws.row_dimensions[start_row + 1].height = 50
            ws.row_dimensions[start_row + 2].height = 80
            ws.row_dimensions[start_row + 3].height = 80

            # 次の週へ
            start_row += 6

    output.seek(0)
    return output.getvalue()